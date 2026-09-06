from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
from PIL import Image

ROOT = Path(r"C:\Omsystems")
OUT = ROOT / "docs" / "client-demo" / "OMSystems_Client_Demo.pptx"
ASSET_DIR = ROOT / "docs" / "manuals" / "screenshots"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(7, 24, 47)
NAVY2 = RGBColor(15, 42, 82)
INK = RGBColor(18, 35, 61)
SLATE = RGBColor(82, 98, 122)
MIST = RGBColor(242, 247, 252)
WHITE = RGBColor(255, 255, 255)
CYAN = RGBColor(24, 183, 196)
BLUE = RGBColor(43, 105, 230)
GREEN = RGBColor(21, 157, 111)
AMBER = RGBColor(244, 166, 35)
RED = RGBColor(218, 76, 92)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def fill(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def rect(slide, x, y, w, h, color, radius=False, line=None, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency
    shape.line.color.rgb = line or color
    shape.line.transparency = 100 if line is None else 0
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def pill(slide, label, x, y, color=CYAN, width=1.45):
    rect(slide, x, y, width, 0.34, color, True)
    text(slide, label.upper(), x, y + 0.055, width, 0.2, 9, WHITE, True, align=PP_ALIGN.CENTER)


def header(slide, eyebrow, title, subtitle=None, dark=False):
    text(slide, eyebrow.upper(), 0.72, 0.42, 4.5, 0.25, 10, CYAN if dark else BLUE, True)
    text(slide, title, 0.7, 0.8, 8.8, 0.62, 28, WHITE if dark else INK, True)
    if subtitle:
        text(slide, subtitle, 0.72, 1.52, 8.8, 0.5, 13, RGBColor(183, 201, 224) if dark else SLATE)


def footer(slide, n, dark=False):
    text(slide, f"OM Systems  /  Security Operations Platform", 0.72, 7.08, 6.3, 0.18, 8, RGBColor(155, 178, 205) if dark else SLATE)
    text(slide, f"{n:02d}", 12.0, 7.02, 0.6, 0.25, 10, RGBColor(155, 178, 205) if dark else SLATE, True, align=PP_ALIGN.RIGHT)


def add_image(slide, filename, x, y, w, h, border=True):
    path = ASSET_DIR / filename
    if not path.exists():
        return False
    with Image.open(path) as im:
        iw, ih = im.size
        ratio = min(w / iw, h / ih)
        nw, nh = iw * ratio, ih * ratio
    px = x + (w - nw) / 2
    py = y + (h - nh) / 2
    if border:
        rect(slide, x - 0.05, y - 0.05, w + 0.1, h + 0.1, WHITE, True, line=RGBColor(215, 225, 237))
    slide.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(nw), height=Inches(nh))
    return True

# Slide 1
s = prs.slides.add_slide(blank); fill(s, NAVY)
rect(s, 8.8, -0.4, 5.2, 8.4, NAVY2, False, transparency=10)
rect(s, 9.65, 0.6, 2.7, 0.08, CYAN)
pill(s, "Client demo", 0.75, 0.7, CYAN, 1.55)
text(s, "OM Systems", 0.75, 1.35, 6.2, 0.55, 32, WHITE, True)
text(s, "Security operations,\nmade visible.", 0.72, 2.0, 7.8, 1.35, 42, WHITE, True)
text(s, "A unified operating layer for live video, AI alerts, incident response, evidence, and fleet health.", 0.78, 3.75, 6.4, 0.75, 17, RGBColor(185, 207, 232))
rect(s, 0.78, 5.35, 5.9, 0.02, RGBColor(63, 91, 126))
text(s, "Prepared for client demonstration", 0.78, 5.65, 4.6, 0.28, 12, RGBColor(185, 207, 232), True)
text(s, "06 September 2026", 0.78, 6.02, 4.6, 0.25, 11, RGBColor(145, 174, 205))
text(s, "OM", 10.0, 2.3, 2.2, 1.0, 64, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(s, "Enterprise Surveillance OS", 9.0, 3.55, 4.0, 0.3, 14, RGBColor(185, 207, 232), True, align=PP_ALIGN.CENTER)
footer(s, 1, True)

# Slide 2
s = prs.slides.add_slide(blank); fill(s, MIST)
header(s, "The operating challenge", "From camera feeds to operational decisions", "Most teams have visibility. The gap is turning visibility into a repeatable response.")
items = [
    ("Fragmented context", "Cameras, alerts, branches, and incidents live in separate views.", BLUE),
    ("Alert overload", "Operators need confidence, prioritization, and suppression of noise.", AMBER),
    ("Slow handoffs", "Evidence and ownership are often assembled after the fact.", RED),
]
for i, (title, body, color) in enumerate(items):
    y = 2.42 + i * 1.22
    rect(s, 0.8, y, 0.12, 0.76, color)
    text(s, title, 1.12, y + 0.02, 3.5, 0.3, 17, INK, True)
    text(s, body, 1.12, y + 0.38, 5.2, 0.34, 12, SLATE)
text(s, "OM Systems brings the workflow together in one operational surface.", 0.82, 6.3, 6.8, 0.35, 16, NAVY, True)
add_image(s, "fresh-command-center.png", 7.25, 2.05, 5.25, 4.35)
footer(s, 2)

# Slide 3
s = prs.slides.add_slide(blank); fill(s, NAVY)
header(s, "One operating surface", "See the whole environment at a glance", "The Command Center gives operators a fast starting point for every shift.", True)
for i, (num, label, color) in enumerate([("01", "Fleet overview", CYAN), ("02", "Live readiness", GREEN), ("03", "Action shortcuts", AMBER)]):
    x = 0.78 + i * 2.25
    text(s, num, x, 2.35, 0.55, 0.35, 19, color, True)
    text(s, label, x, 2.78, 1.75, 0.3, 13, WHITE, True)
    rect(s, x, 3.2, 1.75, 0.03, RGBColor(64, 93, 127))
text(s, "A role-aware workspace for command, triage, and escalation.", 0.8, 5.9, 5.6, 0.4, 16, RGBColor(185, 207, 232), True)
add_image(s, "fresh-command-center.png", 6.85, 1.98, 5.7, 4.7)
footer(s, 3, True)

# Slide 4
s = prs.slides.add_slide(blank); fill(s, MIST)
header(s, "Live monitoring", "Move from overview to live video in one click", "Filter by area, branch, camera status, and operational priority while keeping the fleet in view.")
add_image(s, "fresh-live-video-wall.png", 0.72, 2.1, 7.3, 4.45)
for i, (title, body) in enumerate([
    ("Multi-camera wall", "A single workspace for live streams and camera readiness."),
    ("Operational filters", "Narrow the view by branch, zone, status, or priority."),
    ("Capacity-aware", "Keep operator attention and viewer capacity aligned."),
]):
    y = 2.2 + i * 1.28
    rect(s, 8.55, y, 0.08, 0.72, [BLUE, CYAN, GREEN][i])
    text(s, title, 8.85, y + 0.02, 3.75, 0.28, 16, INK, True)
    text(s, body, 8.85, y + 0.36, 3.55, 0.42, 12, SLATE)
footer(s, 4)

# Slide 5
s = prs.slides.add_slide(blank); fill(s, NAVY)
header(s, "AI intelligence", "Turn detections into prioritized work", "AI Alerts & Incidents brings camera, branch, zone, evidence, and incident context together.", True)
add_image(s, "fresh-ai-alerts.png", 0.72, 2.05, 7.3, 4.6)
for i, (label, color) in enumerate([("Confidence-aware", CYAN), ("Context-rich", BLUE), ("One-click conversion", GREEN)]):
    y = 2.25 + i * 1.18
    rect(s, 8.55, y, 0.12, 0.7, color)
    text(s, label, 8.9, y + 0.02, 3.5, 0.28, 16, WHITE, True)
    text(s, ["Reduce noise before it reaches the operator.", "Understand what happened and where.", "Promote a meaningful alert into a tracked incident."][i], 8.9, y + 0.38, 3.45, 0.38, 12, RGBColor(185, 207, 232))
footer(s, 5, True)

# Slide 6
s = prs.slides.add_slide(blank); fill(s, MIST)
header(s, "Response workflow", "From alert to accountable resolution", "The platform keeps the response chain connected, auditable, and ready for handoff.")
steps = [("01", "Detect", "AI or system signal enters the operating queue.", BLUE), ("02", "Triage", "Operator validates severity, context, and evidence.", AMBER), ("03", "Respond", "Create ownership, incident workflow, and next action.", RED), ("04", "Learn", "Close with outcome, evidence, and audit history.", GREEN)]
for i, (num, title, body, color) in enumerate(steps):
    x = 0.8 + i * 3.05
    rect(s, x, 2.35, 2.45, 2.35, WHITE, True, line=RGBColor(218, 228, 239))
    text(s, num, x + 0.22, 2.62, 0.5, 0.35, 18, color, True)
    text(s, title, x + 0.22, 3.18, 1.9, 0.3, 17, INK, True)
    text(s, body, x + 0.22, 3.7, 1.95, 0.62, 12, SLATE)
    if i < 3:
        text(s, "→", x + 2.58, 3.26, 0.35, 0.35, 22, RGBColor(155, 174, 195), True, align=PP_ALIGN.CENTER)
add_image(s, "fresh-alert-operations.png", 0.9, 5.1, 5.75, 1.25, border=True)
add_image(s, "fresh-security-operations.png", 6.85, 5.1, 5.55, 1.25, border=True)
footer(s, 6)

# Slide 7
s = prs.slides.add_slide(blank); fill(s, NAVY)
header(s, "Trust & governance", "Make every action defensible", "Operational confidence depends on evidence, auditability, and a clear system of record.", True)
points = [
    ("Evidence context", "Snapshots, clips, and alert context travel with the workflow.", CYAN),
    ("Access control", "Role-aware views keep operators focused on their scope.", BLUE),
    ("Fleet health", "Camera, storage, connectivity, and recorder status stay visible.", GREEN),
]
for i, (title, body, color) in enumerate(points):
    y = 2.35 + i * 1.08
    rect(s, 0.85, y, 0.1, 0.68, color)
    text(s, title, 1.18, y, 3.2, 0.28, 16, WHITE, True)
    text(s, body, 1.18, y + 0.35, 4.6, 0.35, 12, RGBColor(185, 207, 232))
add_image(s, "fresh-evidence-vault.png", 6.7, 1.95, 5.8, 4.7)
footer(s, 7, True)

# Slide 8
s = prs.slides.add_slide(blank); fill(s, MIST)
header(s, "Demo close", "A practical path to operational maturity", "Start with visibility. Add intelligence. Standardize response. Scale with confidence.")
for i, (label, detail, color) in enumerate([("See", "Live fleet visibility", BLUE), ("Decide", "Prioritized AI context", CYAN), ("Act", "Connected incident workflow", GREEN)]):
    x = 0.85 + i * 3.98
    rect(s, x, 2.35, 3.25, 1.6, WHITE, True, line=RGBColor(218, 228, 239))
    text(s, str(i + 1).zfill(2), x + 0.22, 2.62, 0.5, 0.3, 18, color, True)
    text(s, label, x + 0.22, 3.06, 1.5, 0.3, 19, INK, True)
    text(s, detail, x + 0.22, 3.47, 2.55, 0.25, 12, SLATE)
rect(s, 0.85, 5.18, 11.6, 0.02, RGBColor(205, 217, 231))
text(s, "Suggested next step", 0.85, 5.55, 2.2, 0.28, 13, BLUE, True)
text(s, "Choose one branch or operational workflow for a focused pilot demo.", 0.85, 5.92, 8.5, 0.4, 20, NAVY, True)
text(s, "Thank you", 10.4, 5.75, 2.0, 0.45, 24, NAVY, True, align=PP_ALIGN.RIGHT)
footer(s, 8)

prs.save(OUT)
print(OUT)
